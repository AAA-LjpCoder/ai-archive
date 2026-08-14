"""文档解析 v2：txt / md / pdf / docx / pptx / xlsx / epub / html / csv → 结构化块
- 标题层级、代码块、表格、列表
- PDF 扫描件 OCR 兜底（PaddleOCR-VL）
- docx 内嵌图片 OCR
"""
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import ZipFile

import mistune

# ---------- 数据结构 ----------


@dataclass
class Block:
    """结构块：heading / paragraph / list / quote / code / table / image_ocr"""
    type: str
    text: str
    level: int = 0          # heading 层级 1-6；非标题为 0
    page_no: int | None = None


@dataclass
class ParsedDoc:
    blocks: list[Block] = field(default_factory=list)
    pages: list[tuple[int, str]] = field(default_factory=list)  # 兼容旧接口 [(page_no, text)]

    @property
    def full_text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


# ---------- 入口 ----------


def parse_file(path: Path, ftype: str) -> ParsedDoc:
    ftype = ftype.lower().lstrip(".")
    if ftype in ("txt", "md", "markdown"):
        return _parse_text(path)
    if ftype == "pdf":
        return _parse_pdf(path)
    if ftype in ("docx", "doc"):
        return _parse_docx(path)
    if ftype == "pptx":
        return _parse_pptx(path)
    if ftype in ("xlsx", "xls"):
        return _parse_xlsx(path)
    if ftype == "epub":
        return _parse_epub(path)
    if ftype in ("html", "htm"):
        return _parse_html(path)
    if ftype == "csv":
        return _parse_csv(path)
    raise ValueError(f"不支持的文件格式: {ftype}")


# ---------- txt / md ----------

_md = mistune.create_markdown(renderer=None, plugins=["table", "strikethrough", "task_lists"])


def _extract_text(node) -> str:
    """递归提取 AST 节点的纯文本"""
    if node is None:
        return ""
    if isinstance(node, list):
        return "".join(_extract_text(n) for n in node)
    if isinstance(node, dict):
        t = node.get("type", "")
        if t == "text":
            return node.get("raw", "") or node.get("text", "")
        if t in ("softbreak", "hardbreak"):
            return "\n"
        if t == "image":
            return node.get("attrs", {}).get("alt", "") or ""
        return _extract_text(node.get("children")) + (node.get("raw", "") or "")
    return ""


def _md_to_blocks(text: str, page_no: int = 1) -> list[Block]:
    """Markdown AST → Block 列表（标题层级 + 结构类型）"""
    blocks: list[Block] = []
    tokens = _md.parse(text)[0]  # mistune v3 返回 (tokens, state)

    def walk(tokens_: list) -> None:
        for tok in tokens_:
            t = tok.get("type", "")
            raw = tok.get("raw", "")
            if t == "heading":
                level = tok.get("attrs", {}).get("level", 1)
                blocks.append(Block(type="heading", text=_extract_text(tok).strip(), level=level, page_no=page_no))
            elif t in ("paragraph", "block_text"):
                body = _extract_text(tok).strip()
                if body:
                    blocks.append(Block(type="paragraph", text=body, page_no=page_no))
            elif t in ("fence", "block_code"):
                code = raw.rstrip("\n")
                if code.strip():
                    blocks.append(Block(type="code", text=code, page_no=page_no))
            elif t == "list":
                walk(tok.get("children", []))
            elif t == "list_item":
                body = _extract_text(tok).strip()
                if body:
                    blocks.append(Block(type="list", text=body, page_no=page_no))
            elif t == "table":
                table_text = _table_to_text(tok)
                if table_text:
                    blocks.append(Block(type="table", text=table_text, page_no=page_no))
            elif t == "block_quote":
                body = _extract_text(tok).strip()
                if body:
                    blocks.append(Block(type="quote", text=body, page_no=page_no))
            elif t == "thematic_break":
                continue
            else:
                body = _extract_text(tok).strip()
                if body and t not in ("html_block",):
                    blocks.append(Block(type="paragraph", text=body, page_no=page_no))

    walk(tokens)
    return blocks


def _table_to_text(tok: dict) -> str:
    """表格 token → 可读文本（表头 + 行，| 分隔）"""
    rows: list[list[str]] = []
    for child in tok.get("children", []):
        t = child.get("type")
        if t == "table_head":
            cells = [_extract_text(c).strip() for c in child.get("children", [])]
            if any(cells):
                rows.append(cells)
        elif t == "table_body":
            for row in child.get("children", []):
                if row.get("type") != "table_row":
                    continue
                cells = [_extract_text(c).strip() for c in row.get("children", [])]
                if any(cells):
                    rows.append(cells)
    if not rows:
        return ""
    return "\n".join(" | ".join(cell for cell in r) for r in rows)


def _parse_text(path: Path) -> ParsedDoc:
    raw = path.read_bytes()
    for enc in ("utf-8", "gbk", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    ftype = path.suffix.lower().lstrip(".")
    if ftype in ("md", "markdown"):
        blocks = _md_to_blocks(text)
    else:
        blocks = _plain_to_blocks(text)

    pages = [(1, text)]
    return ParsedDoc(blocks=blocks, pages=pages)


def _plain_to_blocks(text: str) -> list[Block]:
    """纯文本：启发式识别标题（# / 第X章 / 1.2.3 / 一、）"""
    import re

    blocks: list[Block] = []
    md_heading = re.compile(r"^(#{1,6})\s+(.*)$")
    cn_heading = re.compile(r"^第[一二三四五六七八九十百千\d]+[章节部分卷].*$")
    num_heading = re.compile(r"^(\d+(?:\.\d+)*)[、.．]\s*(.+)$")  # 1.2.3 小节
    cn_num_heading = re.compile(r"^([一二三四五六七八九十]+)、(.+)$")  # 一、二、
    paren_heading = re.compile(r"^（[一二三四五六七八九十\d]+）(.+)$")  # （一）

    def is_heading(line: str) -> int | None:
        m = md_heading.match(line)
        if m:
            return len(m.group(1))
        if cn_heading.match(line) and len(line) <= 30:
            return 1
        m = num_heading.match(line)
        if m and len(line) <= 40:
            return len(m.group(1).split("."))
        if cn_num_heading.match(line) and len(line) <= 30:
            return 1
        if paren_heading.match(line) and len(line) <= 30:
            return 2
        return None

    current: list[str] = []
    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        lvl = is_heading(stripped)
        if lvl is not None:
            if current:
                blocks.append(Block(type="paragraph", text="\n".join(current)))
                current = []
            blocks.append(Block(type="heading", text=stripped, level=lvl))
        else:
            current.append(stripped)
    if current:
        blocks.append(Block(type="paragraph", text="\n".join(current)))
    return blocks


# ---------- pdf ----------


def _parse_pdf(path: Path) -> ParsedDoc:
    import pdfplumber

    doc = ParsedDoc()
    text_pages: list[tuple[int, str]] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            page_texts: list[str] = []
            t = page.extract_text() or ""
            if t.strip():
                page_texts.append(t.strip())
            try:
                for tbl in page.extract_tables():
                    rows = []
                    for r in tbl:
                        cells = [(c or "").replace("\n", " ").strip() for c in r]
                        if any(cells):
                            rows.append(" | ".join(cells))
                    if rows:
                        page_texts.append("\n".join(rows))
            except Exception:  # noqa: BLE001
                pass
            if page_texts:
                text_pages.append((i, "\n\n".join(page_texts)))

    if text_pages:
        for page_no, text in text_pages:
            doc.pages.append((page_no, text))
            for b in _plain_to_blocks(text):
                b.page_no = page_no
                doc.blocks.append(b)
        return doc

    # ---- 扫描件/图片型 PDF：OCR 兜底 ----
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise ValueError("PDF 无可提取文本（扫描件），且服务器未安装 PyMuPDF 无法 OCR") from exc

    from app.services.ocr import ocr_image

    with fitz.open(str(path)) as pdf:
        for i, page in enumerate(pdf, start=1):
            pix = page.get_pixmap(dpi=160)
            img_bytes = pix.tobytes("png")
            text = ocr_image(img_bytes)
            if text.strip():
                page_no = i
                doc.pages.append((page_no, text))
                for b in _plain_to_blocks(text):
                    b.page_no = page_no
                    doc.blocks.append(b)
    if not doc.pages:
        raise ValueError("PDF 解析失败：无文本且 OCR 无结果（可能加密）")
    return doc


# ---------- docx ----------


def _parse_docx(path: Path) -> ParsedDoc:
    import docx

    d = docx.Document(str(path))
    doc = ParsedDoc()
    heading_levels = {
        "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
        "Heading 4": 4, "Heading 5": 5, "Heading 6": 6,
        "标题 1": 1, "标题 2": 2, "标题 3": 3,
        "标题 4": 4, "标题 5": 5, "标题 6": 6,
    }

    def add_paragraph(p) -> None:
        text = (p.text or "").strip()
        if not text:
            return
        style_name = (p.style.name if p.style else "") or ""
        lvl = heading_levels.get(style_name)
        if lvl:
            doc.blocks.append(Block(type="heading", text=text, level=lvl))
        else:
            doc.blocks.append(Block(type="paragraph", text=text))

    def add_table(tbl) -> None:
        rows = []
        for r in tbl.rows:
            cells = [(c.text or "").replace("\n", " ").strip() for c in r.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            doc.blocks.append(Block(type="table", text="\n".join(rows)))

    try:
        for item in d.iter_inner_content():
            if item.__class__.__name__ == "Paragraph":
                add_paragraph(item)
            elif item.__class__.__name__ == "Table":
                add_table(item)
    except Exception:  # noqa: BLE001
        for p in d.paragraphs:
            add_paragraph(p)
        for tbl in d.tables:
            add_table(tbl)

    # ---- 内嵌图片 OCR（补充正文缺失的图片信息） ----
    try:
        _append_docx_image_ocr(path, doc)
    except Exception:  # noqa: BLE001 图片 OCR 失败不阻塞整体
        pass

    if not doc.blocks:
        raise ValueError("docx 无可提取文本")
    return doc


def _append_docx_image_ocr(path: Path, doc: ParsedDoc) -> None:
    """解压 docx 提取 word/media/* 图片 → OCR → 追加为 image_ocr 块（带序号）"""
    from app.services.ocr import ocr_image

    try:
        zf = ZipFile(str(path))
    except Exception:  # noqa: BLE001
        return
    media_names = sorted(
        n for n in zf.namelist()
        if n.startswith("word/media/") and n.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp"))
    )
    for idx, name in enumerate(media_names, start=1):
        data = zf.read(name)
        if len(data) < 4 * 1024:  # 跳过装饰性小图
            continue
        try:
            text = ocr_image(data)
        except Exception:  # noqa: BLE001
            continue
        if text.strip():
            doc.blocks.append(
                Block(type="image_ocr", text=f"【文档图片{idx}】\n{text.strip()}")
            )


# ---------- pptx ----------


def _parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation

    prs = Presentation(str(path))
    doc = ParsedDoc()
    for page_no, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        texts.append(t)
            elif shape.shape_type == 19:  # TABLE
                try:
                    for row in shape.table.rows:
                        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                        if any(cells):
                            texts.append(" | ".join(cells))
                except Exception:  # noqa: BLE001
                    pass
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[备注] {notes}")
        if texts:
            text = "\n".join(texts)
            doc.pages.append((page_no, text))
            for b in _plain_to_blocks(text):
                b.page_no = page_no
                doc.blocks.append(b)
    if not doc.blocks:
        raise ValueError("pptx 无可提取文本")
    return doc


# ---------- xlsx ----------


def _parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    doc = ParsedDoc()
    wb = load_workbook(str(path), read_only=True, data_only=True)
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip().replace("\n", " ") for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            title = f"【工作表】{sheet.title}"
            doc.blocks.append(Block(type="heading", text=title, level=1))
            doc.blocks.append(Block(type="table", text="\n".join(rows)))
    wb.close()
    if not doc.blocks:
        raise ValueError("xlsx 无可提取文本")
    return doc


# ---------- epub ----------


def _parse_epub(path: Path) -> ParsedDoc:
    """epub = zip 里的 XHTML 集合，逐文件解析（复用 HTML 逻辑）"""
    from lxml import etree

    doc = ParsedDoc()
    try:
        zf = ZipFile(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"epub 解压失败: {exc}") from exc

    html_files = sorted(
        n for n in zf.namelist()
        if n.lower().endswith((".xhtml", ".html", ".htm"))
        and not n.startswith("META-INF")
    )
    for idx, name in enumerate(html_files, start=1):
        try:
            content = zf.read(name).decode("utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            continue
        blocks = _html_to_blocks(content, page_no=idx)
        doc.blocks.extend(blocks)
    if not doc.blocks:
        raise ValueError("epub 无可提取文本")
    return doc


# ---------- html ----------


def _html_to_blocks(html: str, page_no: int | None = None) -> list[Block]:
    """HTML → Block 列表：h1-h6 标题、p/li 段落、table 表格、pre/code 代码"""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    # 移除 script/style/nav 噪音
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()

    blocks: list[Block] = []
    body = soup.body or soup

    def handle(el) -> None:
        name = el.name or ""
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            text = el.get_text(" ", strip=True)
            if text:
                blocks.append(Block(type="heading", text=text, level=int(name[1]), page_no=page_no))
        elif name == "table":
            rows = []
            for tr in el.find_all("tr"):
                cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block(type="table", text="\n".join(rows), page_no=page_no))
        elif name in ("pre", "code"):
            text = el.get_text("\n", strip=True)
            if text:
                blocks.append(Block(type="code", text=text, page_no=page_no))
        elif name in ("p", "li", "blockquote", "dd", "dt"):
            text = el.get_text(" ", strip=True)
            if text:
                btype = "quote" if name == "blockquote" else ("list" if name == "li" else "paragraph")
                blocks.append(Block(type=btype, text=text, page_no=page_no))
        else:
            for child in el.children:
                if getattr(child, "name", None):
                    handle(child)

    for child in body.children:
        if getattr(child, "name", None):
            handle(child)
    return blocks


def _parse_html(path: Path) -> ParsedDoc:
    html = path.read_bytes().decode("utf-8", errors="replace")
    blocks = _html_to_blocks(html)
    if not blocks:
        raise ValueError("html 无可提取文本")
    return ParsedDoc(blocks=blocks)


# ---------- csv ----------


def _parse_csv(path: Path) -> ParsedDoc:
    import csv
    import io

    raw = path.read_bytes()
    text = None
    for enc in ("utf-8", "gbk", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    rows: list[str] = []
    try:
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            cells = [(c or "").strip() for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
    except Exception:  # noqa: BLE001 解析失败退回纯文本
        return _parse_text(path)
    if not rows:
        raise ValueError("csv 无可提取文本")
    doc = ParsedDoc()
    doc.blocks.append(Block(type="table", text="\n".join(rows)))
    return doc
