"""解析模块测试"""
from pathlib import Path

import pytest

from app.services.parser import parse_file


def test_parse_txt_utf8(tmp_path: Path):
    p = tmp_path / "a.txt"
    p.write_text("第一行内容\n第二行内容", encoding="utf-8")
    doc = parse_file(p, "txt")
    assert any("第一行内容" in b.text for b in doc.blocks)
    assert any("第二行内容" in b.text for b in doc.blocks)


def test_parse_txt_gbk(tmp_path: Path):
    p = tmp_path / "b.txt"
    p.write_bytes("中文GBK内容测试".encode("gbk"))
    doc = parse_file(p, "txt")
    assert "中文GBK内容测试" in doc.full_text


def test_parse_md_structure(tmp_path: Path):
    """md 应解析出标题层级与正文结构"""
    p = tmp_path / "c.md"
    p.write_text(
        "# 第一章 概述\n\n正文段落。\n\n## 1.1 方法\n\n第二段。\n\n```python\nprint('hi')\n```\n\n| 列A | 列B |\n| --- | --- |\n| 1 | 2 |",
        encoding="utf-8",
    )
    doc = parse_file(p, "md")
    types = [(b.type, b.level) for b in doc.blocks]
    assert ("heading", 1) in types
    assert ("heading", 2) in types
    assert any(b.type == "code" for b in doc.blocks)
    assert any(b.type == "table" for b in doc.blocks)
    # 标题文本不带 # 符号
    assert any(b.type == "heading" and b.text == "第一章 概述" for b in doc.blocks)


def test_parse_md_plain_txt_heading(tmp_path: Path):
    """纯 txt 也应识别常见章节标题"""
    p = tmp_path / "d.txt"
    p.write_text("第一章 概述\n这是内容。\n\n1.2 小节\n更多内容。", encoding="utf-8")
    doc = parse_file(p, "txt")
    headings = [b.text for b in doc.blocks if b.type == "heading"]
    assert "第一章 概述" in headings
    assert "1.2 小节" in headings


def test_parse_unsupported(tmp_path: Path):
    p = tmp_path / "e.xyz"
    p.write_text("x")
    with pytest.raises(ValueError, match="不支持"):
        parse_file(p, "xyz")


# ---------- V2.0 新增格式 ----------


def test_parse_html_structure(tmp_path: Path):
    p = tmp_path / "f.html"
    p.write_text(
        "<html><body><h1>产品手册</h1><h2>第一章</h2>"
        "<p>安装说明。</p><table><tr><th>步骤</th><th>说明</th></tr>"
        "<tr><td>1</td><td>双击</td></tr></table>"
        "<pre><code>print(1)</code></pre></body></html>",
        encoding="utf-8",
    )
    doc = parse_file(p, "html")
    types = [b.type for b in doc.blocks]
    assert "heading" in types and "table" in types and "code" in types
    assert any(b.text == "产品手册" for b in doc.blocks if b.type == "heading")


def test_parse_csv(tmp_path: Path):
    p = tmp_path / "g.csv"
    p.write_text("姓名,部门\n张三,研发", encoding="utf-8")
    doc = parse_file(p, "csv")
    assert doc.blocks[0].type == "table"
    assert "张三" in doc.full_text


def test_parse_xlsx(tmp_path: Path):
    from openpyxl import Workbook

    p = tmp_path / "h.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["月份", "营收"])
    ws.append(["1月", 100])
    wb.save(p)
    doc = parse_file(p, "xlsx")
    assert any("工作表" in b.text for b in doc.blocks)
    assert "1月" in doc.full_text


def test_parse_pptx(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    p = tmp_path / "i.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "季度汇报"
    slide.placeholders[1].text = "营收 500 万"
    prs.save(p)
    doc = parse_file(p, "pptx")
    assert "季度汇报" in doc.full_text
    assert "营收" in doc.full_text


def test_parse_epub(tmp_path: Path):
    import zipfile

    p = tmp_path / "j.epub"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("OEBPS/ch1.xhtml", "<html><body><h1>第一章</h1><p>电子书内容。</p></body></html>")
    doc = parse_file(p, "epub")
    assert any(b.type == "heading" and b.text == "第一章" for b in doc.blocks)
    assert "电子书内容" in doc.full_text
